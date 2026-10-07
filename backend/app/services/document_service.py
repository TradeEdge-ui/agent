"""Document ingestion: text extraction + conservative structured-field extraction.

Design rule from the product spec: never invent values. Every extracted
field either comes from a regex/pattern match against the document's own
text (and is marked with a confidence score) or is left absent. When
confidence is low, `needs_review` is set so the UI can flag it for a human
rather than silently trusting it.
"""
import csv
import io
import json
import re
from pathlib import Path

from app.core.config import get_settings

settings = get_settings()

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".csv", ".png", ".jpg", ".jpeg"}


def extract_text(file_path: str, content_type: str) -> tuple[str, bool]:
    """Returns (text, was_extractable). Never raises on unsupported content."""
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            return text, bool(text.strip())
        if ext == ".docx":
            from docx import Document as DocxDocument

            doc = DocxDocument(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return text, bool(text.strip())
        if ext in (".xlsx", ".xls"):
            from openpyxl import load_workbook

            wb = load_workbook(file_path, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(", ".join(str(c) for c in row if c is not None))
            text = "\n".join(lines)
            return text, bool(text.strip())
        if ext == ".csv":
            with open(file_path, newline="", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                text = "\n".join(", ".join(row) for row in reader)
            return text, bool(text.strip())
        if ext == ".txt":
            with open(file_path, encoding="utf-8", errors="ignore") as f:
                text = f.read()
            return text, bool(text.strip())
        # Images: no OCR engine bundled in this MVP. Flag explicitly instead
        # of pretending we read it.
        return "", False
    except Exception as exc:  # noqa: BLE001 - we want to surface any parser failure as "unreadable"
        return f"[extraction error: {exc}]", False


_MONEY_RE = re.compile(r"(?:RM|MYR|USD|\$|EUR|€|£|GBP)\s?([0-9][0-9,]*\.?[0-9]{0,2})", re.IGNORECASE)
_DATE_RE = re.compile(r"\b(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{2,4}|\d{1,2}\s\w+\s\d{4})\b")
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"\+?\d[\d\s\-()]{7,}\d")
_REF_RE = re.compile(r"\b(?:ref|reference|receipt|invoice|order)\.?\s*#?\s*[:\-]\s*([A-Za-z0-9\-]{4,})", re.IGNORECASE)


def classify_and_extract(text: str) -> dict:
    """Pattern-based structured extraction. Every field carries a confidence
    score derived purely from how specific/unambiguous its match was --
    never a guess at the value itself.
    """
    if not text.strip():
        return {
            "category": "unclassified",
            "fields": {},
            "confidence": 0.0,
            "needs_review": True,
        }

    lower = text.lower()
    amounts = [float(m.replace(",", "")) for m in _MONEY_RE.findall(text)]
    dates = _DATE_RE.findall(text)
    emails = _EMAIL_RE.findall(text)
    phones = _PHONE_RE.findall(text)
    refs = _REF_RE.findall(text)

    category = "other"
    if any(k in lower for k in ["receipt", "paid", "payment received"]):
        category = "receipt"
    elif any(k in lower for k in ["quotation", "quote", "estimate"]):
        category = "quotation"
    elif any(k in lower for k in ["invoice", "bill to"]):
        category = "invoice_like"
    elif emails or phones:
        category = "customer_document"

    fields: dict = {}
    confidences: list[float] = []

    if amounts:
        fields["total_amount"] = {"value": max(amounts), "confidence": 0.7 if len(amounts) == 1 else 0.5}
        confidences.append(fields["total_amount"]["confidence"])
    if dates:
        fields["date"] = {"value": dates[0], "confidence": 0.6}
        confidences.append(0.6)
    if emails:
        fields["email"] = {"value": emails[0], "confidence": 0.9}
        confidences.append(0.9)
    if phones:
        fields["phone"] = {"value": phones[0].strip(), "confidence": 0.5}
        confidences.append(0.5)
    if refs:
        fields["reference_number"] = {"value": refs[0], "confidence": 0.6}
        confidences.append(0.6)

    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    if first_line and category in ("receipt", "quotation", "invoice_like"):
        fields["merchant_or_business_name"] = {"value": first_line[:120], "confidence": 0.3}
        confidences.append(0.3)

    overall_confidence = round(sum(confidences) / len(confidences), 2) if confidences else 0.0
    needs_review = overall_confidence < 0.6 or not fields

    return {
        "category": category,
        "fields": fields,
        "confidence": overall_confidence,
        "needs_review": needs_review,
    }


def fields_to_json(fields: dict) -> str:
    return json.dumps(fields)
